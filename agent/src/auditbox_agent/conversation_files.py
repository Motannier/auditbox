"""Conversation-scoped attachment storage and read-only Agent tools."""

from __future__ import annotations

import hashlib
import json
import re
import shutil
import time
from pathlib import Path
from typing import Any
from uuid import uuid4

from agno.run import RunContext
from agno.tools import Toolkit
from docx import Document as WordDocument

from auditbox_agent.config import Settings
from auditbox_agent.persistence import AgentRepository
from auditbox_agent.storage import (
    StorageValidationError,
    checked_storage_path,
    validate_uploaded_file,
)

SUPPORTED_CONVERSATION_FILE_SUFFIXES = {".md", ".markdown", ".txt", ".docx"}


class ConversationFileError(ValueError):
    """A safe validation error that can be returned by the file API."""


def public_conversation_file(row: dict[str, Any]) -> dict[str, Any]:
    """Return attachment metadata without exposing host filesystem paths."""
    return {
        "id": row["id"],
        "name": row["name"],
        "type": str(row["suffix"]).removeprefix("."),
        "size": int(row["size"]),
        "character_count": int(row["character_count"]),
        "outline": list(row.get("outline") or []),
        "created_at": int(row["created_at"]),
        "status": "ready",
    }


class ConversationFileService:
    """Store originals and normalized text below one conversation-scoped root."""

    def __init__(self, settings: Settings, repository: AgentRepository):
        self.settings = settings
        self.repository = repository
        self.root = settings.conversation_file_directory
        self.root.mkdir(parents=True, exist_ok=True)

    def conversation_root(self, session_id: str) -> Path:
        key = hashlib.sha256(session_id.encode("utf-8")).hexdigest()[:32]
        return self.root / key

    def ingest_staged_file(
        self,
        session_id: str,
        staged_path: Path,
        filename: str,
        content_type: str | None,
    ) -> dict[str, Any]:
        # Browsers normally send a basename, but normalize both path separator
        # styles and control characters before persisting or prompting with it.
        safe_name = Path(filename.replace("\\", "/")).name.strip()
        safe_name = re.sub(r"[\x00-\x1f\x7f]", "_", safe_name)
        suffix = Path(safe_name).suffix.casefold()
        if not safe_name or suffix not in SUPPORTED_CONVERSATION_FILE_SUFFIXES:
            raise ConversationFileError("仅支持 Markdown、TXT 和 DOCX 文档")
        if staged_path.stat().st_size > self.settings.conversation_file_max_bytes:
            raise ConversationFileError(
                f"单个文件不能超过 {self.settings.conversation_file_max_bytes // 1024 // 1024}MB"
            )
        if len(self.repository.list_conversation_files(session_id)) >= (
            self.settings.conversation_file_max_files
        ):
            raise ConversationFileError(
                f"每个对话最多保存 {self.settings.conversation_file_max_files} 个文件"
            )
        try:
            validate_uploaded_file(staged_path, ".md" if suffix == ".markdown" else suffix)
        except StorageValidationError as exc:
            raise ConversationFileError(str(exc)) from exc

        digest = self._sha256(staged_path)
        if self.repository.find_conversation_file_by_hash(session_id, digest):
            raise ConversationFileError("当前对话中已经存在内容相同的文件")
        text = self._extract_text(staged_path, suffix)
        if not text:
            raise ConversationFileError("未能从文件中解析出文本内容")
        if len(text) > self.settings.conversation_file_max_characters:
            raise ConversationFileError(
                f"文档解析后不能超过 {self.settings.conversation_file_max_characters} 个字符"
            )

        file_id = f"file_{uuid4().hex[:16]}"
        folder = self.conversation_root(session_id) / file_id
        original_path = folder / f"original{suffix}"
        text_path = folder / "extracted.md"
        folder.mkdir(parents=True, exist_ok=False)
        try:
            staged_path.replace(original_path)
            text_path.write_text(text, encoding="utf-8")
            row = self.repository.insert_conversation_file(
                {
                    "id": file_id,
                    "session_id": session_id,
                    "name": safe_name,
                    "suffix": suffix,
                    "content_type": content_type,
                    "size": original_path.stat().st_size,
                    "sha256": digest,
                    "original_path": str(original_path.resolve()),
                    "text_path": str(text_path.resolve()),
                    "character_count": len(text),
                    "outline": self._outline(text),
                    "created_at": int(time.time()),
                }
            )
        except Exception:
            shutil.rmtree(folder, ignore_errors=True)
            raise
        return row

    def delete_file(self, session_id: str, file_id: str) -> dict[str, Any] | None:
        row = self.repository.delete_conversation_file(session_id, file_id)
        if row is None:
            return None
        text_path = checked_storage_path(row.get("text_path"), self.root)
        original_path = checked_storage_path(row.get("original_path"), self.root)
        folders = {path.parent for path in (text_path, original_path) if path is not None}
        for folder in folders:
            shutil.rmtree(folder, ignore_errors=True)
        return row

    def delete_conversation(self, session_id: str) -> None:
        self.repository.delete_conversation_files(session_id)
        shutil.rmtree(self.conversation_root(session_id), ignore_errors=True)

    def read_text(self, row: dict[str, Any]) -> str:
        path = checked_storage_path(row.get("text_path"), self.root)
        if path is None:
            raise ConversationFileError("文件正文不存在或存储路径无效")
        return path.read_text(encoding="utf-8")

    def question_excerpts(self, session_id: str, question: str) -> list[dict[str, Any]]:
        """Supply bounded attachment evidence for Q&A without model tool calls."""
        rows = self.repository.list_conversation_files(session_id)
        named = [row for row in rows if str(row["name"]).casefold() in question.casefold()]
        if not named and not re.search(r"附件|文件|文档|资料|手册|材料", question):
            return []
        terms = ConversationFileTools._query_terms(question)
        candidates: list[dict[str, Any]] = []
        for row in named or rows:
            try:
                text = self.read_text(row)
            except (ConversationFileError, OSError):
                candidates.append({"file_id": row["id"], "name": row["name"],
                                   "error": "附件正文当前不可用"})
                continue
            ranked: list[tuple[int, int, str]] = []
            for start, value in ConversationFileTools._windows(text):
                folded = value.casefold()
                ranked.append((sum(folded.count(term) for term in terms), start, value))
            ranked.sort(key=lambda item: item[0], reverse=True)
            for score, start, value in ranked[:2]:
                candidates.append({
                    "file_id": row["id"], "name": row["name"], "start": start,
                    "total_characters": len(text), "score": score, "content": value,
                })
        candidates.sort(key=lambda item: (-int(item.get("score", 0)), str(item["file_id"])))
        remaining = min(
            self.settings.qa_file_context_characters,
            self.settings.context_max_input_tokens // 4,
        )
        excerpts: list[dict[str, Any]] = []
        for candidate in candidates[:8]:
            candidate.pop("score", None)
            if "error" in candidate:
                excerpts.append(candidate)
                continue
            if remaining <= 0:
                break
            content = candidate["content"][:remaining]
            candidate["content"] = content
            candidate["end"] = candidate["start"] + len(content)
            candidate["partial"] = len(content) < candidate["total_characters"]
            remaining -= len(content)
            excerpts.append(candidate)
        return excerpts

    @staticmethod
    def _sha256(path: Path) -> str:
        digest = hashlib.sha256()
        with path.open("rb") as source:
            while chunk := source.read(1024 * 1024):
                digest.update(chunk)
        return digest.hexdigest()

    @staticmethod
    def _normalize(value: str) -> str:
        value = value.replace("\x00", "").replace("\r\n", "\n").replace("\r", "\n")
        value = re.sub(r"[ \t]+\n", "\n", value)
        return re.sub(r"\n{4,}", "\n\n\n", value).strip()

    def _extract_text(self, path: Path, suffix: str) -> str:
        if suffix != ".docx":
            return self._normalize(path.read_text(encoding="utf-8"))
        document = WordDocument(str(path))
        parts: list[str] = []
        for paragraph in document.paragraphs:
            value = paragraph.text.strip()
            if not value:
                continue
            style = str(paragraph.style.name if paragraph.style is not None else "").casefold()
            match = re.search(r"heading\s*([1-6])", style)
            parts.append(f"{'#' * int(match.group(1))} {value}" if match else value)
        for table in document.tables:
            rows = [
                " | ".join(cell.text.strip().replace("\n", " ") for cell in row.cells)
                for row in table.rows
            ]
            if rows:
                parts.extend(rows)
        return self._normalize("\n\n".join(parts))

    @staticmethod
    def _outline(text: str) -> list[str]:
        headings = [
            match.group(2).strip()
            for line in text.splitlines()
            if (match := re.match(r"^(#{1,6})\s+(.+?)\s*$", line))
        ]
        return headings[:50]


class ConversationFileTools(Toolkit):
    """Small, read-only tools scoped by Agno's injected session id."""

    def __init__(self, service: ConversationFileService):
        self.service = service
        super().__init__(
            name="conversation_files",
            tools=[
                self.list_conversation_files,
                self.get_conversation_file_outline,
                self.search_conversation_files,
                self.read_conversation_file,
            ],
        )

    def list_conversation_files(self, run_context: RunContext) -> str:
        """列出当前对话文件区中可读取的文件及其 file_id。"""
        rows = self.service.repository.list_conversation_files(self._session_id(run_context))
        return json.dumps(
            [public_conversation_file(row) for row in rows], ensure_ascii=False
        )

    def get_conversation_file_outline(self, run_context: RunContext, file_id: str) -> str:
        """读取指定对话文件的标题目录；先用 file_id，不要把文件名当作系统路径。"""
        row = self._file(run_context, file_id)
        return json.dumps(
            {
                "file_id": row["id"],
                "name": row["name"],
                "character_count": row["character_count"],
                "outline": row.get("outline") or [],
            },
            ensure_ascii=False,
        )

    def search_conversation_files(
        self,
        run_context: RunContext,
        query: str,
        file_id: str = "",
        max_results: int = 5,
    ) -> str:
        """按关键词搜索当前对话文件，返回少量相关原文片段。file_id 留空时搜索全部文件。"""
        query = query.strip()
        if not query:
            return json.dumps({"error": "搜索词不能为空"}, ensure_ascii=False)
        rows = (
            [self._file(run_context, file_id)]
            if file_id
            else self.service.repository.list_conversation_files(self._session_id(run_context))
        )
        terms = self._query_terms(query)
        matches: list[dict[str, Any]] = []
        for row in rows:
            text = self.service.read_text(row)
            for start, value in self._windows(text):
                folded = value.casefold()
                score = sum(folded.count(term) * max(1, len(term)) for term in terms)
                if query.casefold() in folded:
                    score += 100
                if score:
                    matches.append(
                        {
                            "file_id": row["id"],
                            "name": row["name"],
                            "start": start,
                            "score": score,
                            "content": value[:1200],
                        }
                    )
        matches.sort(
            key=lambda item: (
                -int(item["score"]),
                str(item["file_id"]),
                int(item["start"]),
            )
        )
        return json.dumps(
            {"query": query, "results": matches[: max(1, min(max_results, 8))]},
            ensure_ascii=False,
        )

    def read_conversation_file(
        self,
        run_context: RunContext,
        file_id: str,
        start: int = 0,
        length: int = 4000,
    ) -> str:
        """按字符范围读取一个对话文件。长文档应分段读取，单次最多返回 6000 字符。"""
        row = self._file(run_context, file_id)
        text = self.service.read_text(row)
        start = max(0, min(start, len(text)))
        limit = max(200, min(length, 6000, self.service.settings.conversation_file_read_characters))
        end = min(len(text), start + limit)
        return json.dumps(
            {
                "file_id": row["id"],
                "name": row["name"],
                "start": start,
                "end": end,
                "total_characters": len(text),
                "has_more": end < len(text),
                "content": text[start:end],
            },
            ensure_ascii=False,
        )

    def _file(self, run_context: RunContext, file_id: str) -> dict[str, Any]:
        row = self.service.repository.get_conversation_file(
            self._session_id(run_context), file_id.strip()
        )
        if row is None:
            raise ConversationFileError("当前对话文件区中不存在该 file_id")
        return row

    @staticmethod
    def _session_id(run_context: RunContext) -> str:
        session_id = str(run_context.session_id or "")
        if not session_id:
            raise ConversationFileError("当前工具调用缺少对话标识")
        return session_id

    @staticmethod
    def _query_terms(query: str) -> list[str]:
        folded = query.casefold()
        terms = {value for value in re.findall(r"[a-z0-9_.:/-]{2,}", folded)}
        for sequence in re.findall(r"[\u3400-\u9fff]+", folded):
            if len(sequence) <= 12:
                terms.add(sequence)
            terms.update(sequence[index : index + 2] for index in range(len(sequence) - 1))
        return sorted(terms, key=lambda value: (-len(value), value))[:40] or [folded]

    @staticmethod
    def _windows(text: str, size: int = 1400, overlap: int = 160) -> list[tuple[int, str]]:
        values: list[tuple[int, str]] = []
        start = 0
        while start < len(text):
            end = min(len(text), start + size)
            if end < len(text):
                natural = max(text.rfind("\n\n", start, end), text.rfind("\n", start, end))
                if natural > start + size // 2:
                    end = natural
            raw = text[start:end]
            value = raw.strip()
            if value:
                values.append((start + len(raw) - len(raw.lstrip()), value))
            if end >= len(text):
                break
            start = max(start + 1, end - overlap)
        return values
