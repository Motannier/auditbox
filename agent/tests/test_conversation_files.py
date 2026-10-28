import json
from io import BytesIO

import pytest
from agno.run import RunContext
from docx import Document
from fastapi.testclient import TestClient

from auditbox_agent.conversation_files import ConversationFileError
from auditbox_agent.runtime import AgentRuntime


def _create_conversation(client: TestClient, title: str = "文件区测试") -> str:
    response = client.post("/api/v1/conversations", json={"title": title})
    assert response.status_code == 201
    return str(response.json()["session_id"])


def test_conversation_file_lifecycle_and_scoped_agent_read(
    client: TestClient,
    runtime: AgentRuntime,
) -> None:
    conversation_id = _create_conversation(client)
    uploaded = client.post(
        f"/api/v1/conversations/{conversation_id}/files",
        files={
            "file": (
                "../网络排障.md",
                "# 网络排障\n\n唯一标记：WORKSPACE-42\n\n检查接口和路由。".encode(),
                "text/markdown",
            )
        },
    )
    assert uploaded.status_code == 201
    file_data = uploaded.json()
    assert file_data["name"] == "网络排障.md"
    assert file_data["outline"] == ["网络排障"]
    assert "content" not in file_data
    assert "path" not in file_data

    listed = client.get(f"/api/v1/conversations/{conversation_id}/files")
    assert listed.json()["data"] == [file_data]

    context = RunContext(
        run_id="file-tool-run",
        session_id=conversation_id,
        user_id=runtime.settings.local_user_id,
    )
    search = json.loads(
        runtime.conversation_file_tools.search_conversation_files(context, "唯一标记")
    )
    assert search["results"][0]["file_id"] == file_data["id"]
    assert "WORKSPACE-42" in search["results"][0]["content"]
    read = json.loads(
        runtime.conversation_file_tools.read_conversation_file(
            context, file_data["id"], start=0, length=400
        )
    )
    assert "WORKSPACE-42" in read["content"]
    assert read["has_more"] is False

    other_id = _create_conversation(client, "另一个文件区")
    with pytest.raises(ConversationFileError, match="不存在"):
        runtime.conversation_file_tools.read_conversation_file(
            RunContext(run_id="other-run", session_id=other_id),
            file_data["id"],
        )

    deleted = client.delete(
        f"/api/v1/conversations/{conversation_id}/files/{file_data['id']}"
    )
    assert deleted.status_code == 204
    assert client.get(f"/api/v1/conversations/{conversation_id}/files").json() == {
        "data": []
    }


def test_docx_is_normalized_and_conversation_delete_cleans_workspace(
    client: TestClient,
    runtime: AgentRuntime,
) -> None:
    document = Document()
    document.add_heading("设备交接", level=1)
    document.add_paragraph("交换机管理地址为 192.0.2.10。")
    content = BytesIO()
    document.save(content)

    conversation_id = _create_conversation(client)
    uploaded = client.post(
        f"/api/v1/conversations/{conversation_id}/files",
        files={
            "file": (
                "设备交接.docx",
                content.getvalue(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    assert uploaded.status_code == 201
    assert uploaded.json()["outline"] == ["设备交接"]
    row = runtime.repository.get_conversation_file(
        conversation_id, uploaded.json()["id"]
    )
    assert row is not None
    assert "192.0.2.10" in runtime.conversation_files.read_text(row)
    conversation_root = runtime.conversation_files.conversation_root(conversation_id)
    assert conversation_root.is_dir()

    assert client.delete(f"/api/v1/conversations/{conversation_id}").status_code == 204
    assert not conversation_root.exists()


def test_conversation_file_rejects_duplicates_and_unsupported_types(
    client: TestClient,
) -> None:
    conversation_id = _create_conversation(client)
    payload = {
        "file": ("notes.txt", "同一份内容".encode(), "text/plain")
    }
    assert (
        client.post(
            f"/api/v1/conversations/{conversation_id}/files", files=payload
        ).status_code
        == 201
    )
    duplicate = client.post(
        f"/api/v1/conversations/{conversation_id}/files", files=payload
    )
    assert duplicate.status_code == 422
    assert "内容相同" in duplicate.json()["detail"]
    unsupported = client.post(
        f"/api/v1/conversations/{conversation_id}/files",
        files={"file": ("script.py", b"print('no')", "text/x-python")},
    )
    assert unsupported.status_code == 422
    assert "仅支持" in unsupported.json()["detail"]
